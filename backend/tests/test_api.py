"""Integration tests for Evidence-First Resume Triage API.

Covers:
- Server-root GET / serving static frontend HTML
- Every contract route from docs/CONTRACT.md
- PDF, DOCX, and TXT resume uploads with fact extraction
- Frontend-facing response shapes and contract compliance
- Score changes after human evidence accept/reject decisions
- Input validation and 404/422 error handling
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import uuid
from pathlib import Path

import pytest

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"

# Isolated throwaway database for test run
os.environ["DB_PATH"] = str(Path(tempfile.gettempdir()) / f"resume_triage_test_{uuid.uuid4().hex}.db")

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402

client = TestClient(app)

JOB_DESCRIPTION = """Senior Backend Engineer

Must-Have Requirements:
- At least 5 years of professional Python experience
- Experience with FastAPI

Nice-to-Have Requirements:
- Experience with Kubernetes
- Familiarity with AWS
"""

FIXTURE_FILES = ["strong_match.txt", "junior_missing.txt", "contradicted_signal.txt"]


# ---------------------------------------------------------------------------
# Test Helpers: Document Generation & Contract Shape Assertions
# ---------------------------------------------------------------------------

def _generate_pdf_bytes(name: str, email: str, body: str) -> bytes:
    """Generate in-memory PDF resume bytes using PyMuPDF."""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open()
    page = doc.new_page()
    content = f"{name}\n{email}\n\n{body}"
    page.insert_text((50, 72), content)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _generate_docx_bytes(name: str, email: str, body: str) -> bytes:
    """Generate in-memory DOCX resume bytes using python-docx."""
    import docx
    doc = docx.Document()
    doc.add_paragraph(name)
    doc.add_paragraph(email)
    doc.add_paragraph(body)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _assert_requirement_shape(req: dict) -> None:
    """Assert a requirement dictionary adheres to CONTRACT.md."""
    assert isinstance(req, dict)
    assert isinstance(req.get("id"), str) and req["id"]
    assert isinstance(req.get("text"), str) and req["text"]
    assert req.get("type") in {"must", "nice"}


def _assert_evidence_shape(ev: dict) -> None:
    """Assert an evidence dictionary adheres to CONTRACT.md."""
    assert isinstance(ev, dict)
    assert isinstance(ev.get("snippet"), str)
    assert isinstance(ev.get("confidence"), (int, float))
    assert 0.0 <= ev["confidence"] <= 1.0
    assert ev.get("status") in {"verified", "uncertain", "contradicted", "missing"}


def _assert_match_shape(match: dict) -> None:
    """Assert a requirement match dictionary adheres to CONTRACT.md."""
    assert isinstance(match, dict)
    assert isinstance(match.get("requirement_id"), str) and match["requirement_id"]
    assert isinstance(match.get("evidence"), list)
    for ev in match["evidence"]:
        _assert_evidence_shape(ev)
    assert match.get("status") in {"verified", "uncertain", "contradicted", "missing"}


def _assert_candidate_shape(candidate: dict) -> None:
    """Assert a candidate dictionary adheres to CONTRACT.md."""
    assert isinstance(candidate, dict)
    assert isinstance(candidate.get("id"), str) and candidate["id"]
    assert isinstance(candidate.get("job_id"), str) and candidate["job_id"]
    assert isinstance(candidate.get("name"), str)
    assert isinstance(candidate.get("email"), str)
    assert isinstance(candidate.get("skills"), list)
    for skill in candidate["skills"]:
        assert isinstance(skill, str)
    assert isinstance(candidate.get("experience_years"), (int, float))
    assert candidate["experience_years"] >= 0
    assert isinstance(candidate.get("raw_text"), str)
    assert isinstance(candidate.get("matches"), list)
    for match in candidate["matches"]:
        _assert_match_shape(match)
    assert isinstance(candidate.get("score"), (int, float))
    assert 0.0 <= candidate["score"] <= 1.0


def _create_job(description: str = JOB_DESCRIPTION) -> dict:
    response = client.post("/api/jobs", json={"description": description})
    assert response.status_code == 200
    data = response.json()
    assert "job_id" in data
    assert "requirements" in data
    for req in data["requirements"]:
        _assert_requirement_shape(req)
    return data


# ---------------------------------------------------------------------------
# Test Cases
# ---------------------------------------------------------------------------

def test_server_root_serves_frontend():
    """Server-root GET / and GET /dashboard.html must serve static frontend HTML."""
    root_response = client.get("/")
    assert root_response.status_code == 200
    assert "text/html" in root_response.headers.get("content-type", "")
    assert len(root_response.text) > 0
    # Must contain HTML content
    assert "<html" in root_response.text.lower() or "<!doctype html" in root_response.text.lower()

    dashboard_response = client.get("/dashboard.html")
    assert dashboard_response.status_code == 200
    assert "text/html" in dashboard_response.headers.get("content-type", "")
    assert len(dashboard_response.text) > 0


def test_contract_response_shapes_and_every_contract_route():
    """Verify every route in docs/CONTRACT.md adheres to the required response shapes."""
    # 1. POST /api/jobs -> {job_id:str, requirements:[Requirement]}
    job = _create_job()
    job_id = job["job_id"]
    assert isinstance(job_id, str) and job_id
    assert len(job["requirements"]) >= 2
    for req in job["requirements"]:
        _assert_requirement_shape(req)

    # 1b. GET /api/jobs/{job_id} -> {id:str, description:str, requirements:[Requirement]}
    job_get_response = client.get(f"/api/jobs/{job_id}")
    assert job_get_response.status_code == 200
    job_data = job_get_response.json()
    assert job_data["id"] == job_id
    assert "description" in job_data
    for req in job_data["requirements"]:
        _assert_requirement_shape(req)

    # 2. POST /api/candidates/upload multipart -> {candidates:[Candidate]}
    files = [
        ("files", ("dana.txt", (FIXTURES_DIR / "strong_match.txt").read_bytes(), "text/plain"))
    ]
    upload_response = client.post(
        "/api/candidates/upload",
        data={"job_id": job_id},
        files=files,
    )
    assert upload_response.status_code == 200
    upload_data = upload_response.json()
    assert "candidates" in upload_data
    assert len(upload_data["candidates"]) == 1
    candidate = upload_data["candidates"][0]
    _assert_candidate_shape(candidate)
    candidate_id = candidate["id"]

    # 3. GET /api/jobs/{job_id}/candidates -> {candidates:[Candidate]} sorted by score desc
    list_response = client.get(f"/api/jobs/{job_id}/candidates")
    assert list_response.status_code == 200
    list_data = list_response.json()
    assert "candidates" in list_data
    for cand in list_data["candidates"]:
        _assert_candidate_shape(cand)

    # 4. PATCH /api/candidates/{id}/facts -> {candidate:Candidate}
    patch_response = client.patch(
        f"/api/candidates/{candidate_id}/facts",
        json={"field": "experience_years", "value": 7.5},
    )
    assert patch_response.status_code == 200
    patch_data = patch_response.json()
    assert "candidate" in patch_data
    _assert_candidate_shape(patch_data["candidate"])
    assert patch_data["candidate"]["experience_years"] == 7.5

    # 5. POST /api/candidates/{id}/evidence/{requirement_id}/status -> {candidate:Candidate}
    target_req_id = job["requirements"][0]["id"]
    status_response = client.post(
        f"/api/candidates/{candidate_id}/evidence/{target_req_id}/status",
        json={"status": "verified"},
    )
    assert status_response.status_code == 200
    status_data = status_response.json()
    assert "candidate" in status_data
    _assert_candidate_shape(status_data["candidate"])

    # 6. GET /api/jobs/{job_id}/export -> {markdown:str}
    export_response = client.get(f"/api/jobs/{job_id}/export")
    assert export_response.status_code == 200
    export_data = export_response.json()
    assert "markdown" in export_data
    assert isinstance(export_data["markdown"], str)
    assert len(export_data["markdown"]) > 0


def test_pdf_and_docx_uploads():
    """Upload multipart .pdf and .docx resumes and verify fact extraction and scoring."""
    job = _create_job()
    job_id = job["job_id"]

    pdf_bytes = _generate_pdf_bytes(
        name="Alex Morgan",
        email="alex.morgan@example.com",
        body="Senior Backend Engineer with 6 years of professional experience in Python and FastAPI. "
             "Skills: Python, FastAPI, Docker",
    )

    docx_bytes = _generate_docx_bytes(
        name="Jordan Taylor",
        email="jordan.taylor@example.com",
        body="Infrastructure Engineer with 4 years of experience in Kubernetes and AWS cloud deployments. "
             "Skills: Kubernetes, AWS, Python",
    )

    upload_response = client.post(
        "/api/candidates/upload",
        data={"job_id": job_id},
        files=[
            ("files", ("alex_resume.pdf", pdf_bytes, "application/pdf")),
            ("files", ("jordan_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
    )
    assert upload_response.status_code == 200
    candidates = upload_response.json()["candidates"]
    assert len(candidates) == 2

    by_name = {c["name"]: c for c in candidates}
    assert "Alex Morgan" in by_name
    assert "Jordan Taylor" in by_name

    alex = by_name["Alex Morgan"]
    _assert_candidate_shape(alex)
    assert alex["email"] == "alex.morgan@example.com"
    assert "Python" in alex["skills"]
    assert "FastAPI" in alex["skills"]
    assert alex["experience_years"] >= 5.0
    assert alex["score"] > 0.0

    jordan = by_name["Jordan Taylor"]
    _assert_candidate_shape(jordan)
    assert jordan["email"] == "jordan.taylor@example.com"
    assert "Kubernetes" in jordan["skills"]
    assert "AWS" in jordan["skills"]
    assert jordan["experience_years"] >= 4.0

    # Verify both candidates are retrievable from GET /api/jobs/{job_id}/candidates
    list_response = client.get(f"/api/jobs/{job_id}/candidates")
    assert list_response.status_code == 200
    retrieved = list_response.json()["candidates"]
    retrieved_names = {c["name"] for c in retrieved}
    assert "Alex Morgan" in retrieved_names
    assert "Jordan Taylor" in retrieved_names


def test_score_changes_after_evidence_accept_and_reject():
    """Verify candidate score increases when evidence is accepted and decreases when rejected."""
    job = _create_job(
        """Backend Developer
Must-Have Requirements:
- At least 5 years of Python experience
- Experience with FastAPI

Nice-to-Have Requirements:
- Experience with Kubernetes
"""
    )
    job_id = job["job_id"]
    must_reqs = [r for r in job["requirements"] if r["type"] == "must"]
    nice_reqs = [r for r in job["requirements"] if r["type"] == "nice"]
    assert len(must_reqs) == 2
    assert len(nice_reqs) == 1

    must_id_1 = must_reqs[0]["id"]
    must_id_2 = must_reqs[1]["id"]
    nice_id = nice_reqs[0]["id"]

    # Upload a candidate with no matching skills or experience initially
    upload_res = client.post(
        "/api/candidates/upload",
        data={"job_id": job_id},
        files=[
            (
                "files",
                (
                    "candidate.txt",
                    b"Pat Quinn\npat@example.com\nEntry-level developer with 0 years of experience.",
                    "text/plain",
                ),
            )
        ],
    )
    assert upload_res.status_code == 200
    cand = upload_res.json()["candidates"][0]
    cand_id = cand["id"]
    initial_score = cand["score"]
    assert initial_score == 0.0

    # 1. Accept first must-have requirement (score must increase)
    accept_1 = client.post(
        f"/api/candidates/{cand_id}/evidence/{must_id_1}/status",
        json={"status": "verified"},
    )
    assert accept_1.status_code == 200
    cand_1 = accept_1.json()["candidate"]
    score_after_must1 = cand_1["score"]
    assert score_after_must1 > initial_score
    # Must-have weight is 2; total possible is 2 + 2 + 1 = 5; 2/5 = 0.4
    assert pytest.approx(score_after_must1, abs=1e-3) == 2 / 5

    # 2. Accept second must-have requirement (score must increase further)
    accept_2 = client.post(
        f"/api/candidates/{cand_id}/evidence/{must_id_2}/status",
        json={"status": "verified"},
    )
    assert accept_2.status_code == 200
    cand_2 = accept_2.json()["candidate"]
    score_after_must2 = cand_2["score"]
    assert score_after_must2 > score_after_must1
    # 4/5 = 0.8
    assert pytest.approx(score_after_must2, abs=1e-3) == 4 / 5

    # 3. Accept nice-to-have requirement (score reaches maximum 1.0)
    accept_3 = client.post(
        f"/api/candidates/{cand_id}/evidence/{nice_id}/status",
        json={"status": "verified"},
    )
    assert accept_3.status_code == 200
    cand_3 = accept_3.json()["candidate"]
    score_after_nice = cand_3["score"]
    assert score_after_nice > score_after_must2
    assert score_after_nice == 1.0

    # 4. Reject must-have requirement as contradicted (score must drop)
    reject_must = client.post(
        f"/api/candidates/{cand_id}/evidence/{must_id_1}/status",
        json={"status": "contradicted"},
    )
    assert reject_must.status_code == 200
    cand_rejected = reject_must.json()["candidate"]
    score_after_reject = cand_rejected["score"]
    assert score_after_reject < 1.0
    # Remaining: must_2 (2) + nice (1) = 3/5 = 0.6
    assert pytest.approx(score_after_reject, abs=1e-3) == 3 / 5

    # 5. Reject nice-to-have as missing (score must drop further)
    reject_nice = client.post(
        f"/api/candidates/{cand_id}/evidence/{nice_id}/status",
        json={"status": "missing"},
    )
    assert reject_nice.status_code == 200
    cand_rejected_2 = reject_nice.json()["candidate"]
    assert cand_rejected_2["score"] < score_after_reject
    # Remaining: must_2 (2) = 2/5 = 0.4
    assert pytest.approx(cand_rejected_2["score"], abs=1e-3) == 2 / 5


def test_full_upload_to_export_flow():
    """Existing full end-to-end integration test with shape validations."""
    job = _create_job()
    job_id = job["job_id"]
    assert job_id
    assert any(requirement["type"] == "must" for requirement in job["requirements"])
    assert any(requirement["type"] == "nice" for requirement in job["requirements"])

    files = [
        ("files", (name, (FIXTURES_DIR / name).read_bytes(), "text/plain"))
        for name in FIXTURE_FILES
    ]
    upload_response = client.post(
        "/api/candidates/upload",
        data={"job_id": job_id},
        files=files,
    )
    assert upload_response.status_code == 200
    candidates = upload_response.json()["candidates"]
    assert len(candidates) == len(FIXTURE_FILES)
    for candidate in candidates:
        _assert_candidate_shape(candidate)
        assert candidate["job_id"] == job_id
        assert candidate["matches"]
        assert 0.0 <= candidate["score"] <= 1.0

    # Fetch candidates: the endpoint must rank them by score, descending.
    list_response = client.get(f"/api/jobs/{job_id}/candidates")
    assert list_response.status_code == 200
    ranked = list_response.json()["candidates"]
    assert len(ranked) == len(FIXTURE_FILES)
    scores = [candidate["score"] for candidate in ranked]
    assert scores == sorted(scores, reverse=True)

    names = {candidate["name"]: candidate for candidate in ranked}
    assert "Dana Reyes" in names
    assert "Jamie Lee" in names
    assert names["Dana Reyes"]["score"] > names["Jamie Lee"]["score"]

    # Correct a fact and confirm score recalculation
    jamie = names["Jamie Lee"]
    patch_response = client.patch(
        f"/api/candidates/{jamie['id']}/facts",
        json={"field": "experience_years", "value": 6},
    )
    assert patch_response.status_code == 200
    corrected = patch_response.json()["candidate"]
    _assert_candidate_shape(corrected)
    assert corrected["experience_years"] == 6
    assert corrected["score"] > jamie["score"]

    # Export report
    export_response = client.get(f"/api/jobs/{job_id}/export")
    assert export_response.status_code == 200
    markdown = export_response.json()["markdown"]
    assert "Dana Reyes" in markdown
    assert "Jamie Lee" in markdown
    assert "Sam Okafor" in markdown
    assert "decision support" in markdown.lower()
    for requirement in job["requirements"]:
        if requirement["type"] == "must":
            assert requirement["text"] in markdown


def test_error_handling_and_validation_across_routes():
    """Verify error responses across all contract endpoints."""
    # 1. POST /api/jobs with empty description -> 400 or 422
    empty_job_response = client.post("/api/jobs", json={"description": "   "})
    assert empty_job_response.status_code in {400, 422}

    # 2. GET /api/jobs/missing_id -> 404
    missing_job_response = client.get("/api/jobs/job_nonexistent_12345")
    assert missing_job_response.status_code == 404

    # 3. GET /api/jobs/missing_id/export -> 404
    export_missing = client.get("/api/jobs/job_nonexistent_12345/export")
    assert export_missing.status_code == 404

    # 4. PATCH /api/candidates/missing_id/facts -> 404
    patch_missing = client.patch(
        "/api/candidates/cand_missing_12345/facts",
        json={"field": "experience_years", "value": 5},
    )
    assert patch_missing.status_code == 404

    # Create a valid candidate for field validation checks
    job = _create_job()
    upload = client.post(
        "/api/candidates/upload",
        data={"job_id": job["job_id"]},
        files=[("files", ("test.txt", b"Jane Doe\njane@example.com\n3 years Python", "text/plain"))],
    )
    cand_id = upload.json()["candidates"][0]["id"]
    req_id = job["requirements"][0]["id"]

    # 5. PATCH with invalid field name -> 422
    patch_invalid_field = client.patch(
        f"/api/candidates/{cand_id}/facts",
        json={"field": "salary", "value": 100000},
    )
    assert patch_invalid_field.status_code == 422

    # 6. PATCH with negative experience -> 422
    patch_negative_exp = client.patch(
        f"/api/candidates/{cand_id}/facts",
        json={"field": "experience_years", "value": -2.0},
    )
    assert patch_negative_exp.status_code == 422

    # 7. POST evidence status on unknown candidate -> 404
    evidence_missing_cand = client.post(
        f"/api/candidates/cand_missing_12345/evidence/{req_id}/status",
        json={"status": "verified"},
    )
    assert evidence_missing_cand.status_code == 404

    # 8. POST evidence status on unknown requirement -> 404
    evidence_missing_req = client.post(
        f"/api/candidates/{cand_id}/evidence/req_nonexistent_9999/status",
        json={"status": "verified"},
    )
    assert evidence_missing_req.status_code == 404

    # 9. POST evidence status with invalid status -> 422
    evidence_invalid_status = client.post(
        f"/api/candidates/{cand_id}/evidence/{req_id}/status",
        json={"status": "somewhat_acceptable"},
    )
    assert evidence_invalid_status.status_code == 422
